"""Lógica de generación de la Base de Gestión por Campaña de Trabajo.

Implementa la especificación técnica "Generación Automática de Base de Cartera
por Campaña de Trabajo". Este módulo no depende de Streamlit para poder
probarse y reutilizarse de forma independiente.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd

# --------------------------------------------------------------------------
# Constantes
# --------------------------------------------------------------------------

COLUMNAS_SALIDA = [
    "ZONA",
    "REGION",
    "RUTA",
    "DIVISION",
    "ID COBRADOR",
    "NoDama",
    "Direccion",
    "Direccion Calle",
    "Colonia",
    "Municipio / Poblacion",
    "Cp",
    "Estado",
    "Zona (Urbano/Rural)",
    "Referencia",
    "AnioSaldo",
    "CampaniaSaldo",
    "DigitoVerificador",
    "Concatenado",
    "Fecha de cierre",
    "Morosidad",
    "Campaña de trabajo",
    "Referencia de Pago",
]

COLUMNAS_BASE = COLUMNAS_SALIDA + ["requiere_revision", "motivo_revision", "_fila_excel"]

COLUMNAS_CARTERA = [
    "ZONA",
    "NoDama",
    "Direccion",
    "Referencia",
    "AnioSaldo",
    "CampaniaSaldo",
    "DigitoVerificador",
]

RE_CP = re.compile(r"CP\s*(\d{1,5})\s*$", re.IGNORECASE)
RE_CP_FINAL = re.compile(r"\s*CP\s*\d*\s*$", re.IGNORECASE)
PALABRAS_CLAVE = ("No", "Mza", "Int", "Lt")
RE_CLAVE = re.compile(r"\b(No|Mza|Int|Lt)\b\.?", re.IGNORECASE)
RE_TOKEN = re.compile(r"\s*([0-9A-Za-zÀ-ÿ]+(?:[-/][0-9A-Za-zÀ-ÿ]+)*)")

REGLAS_TIPO_ASENTAMIENTO = {
    "fraccionamiento": "Urbano",
    "colonia": "Urbano",
    "unidad habitacional": "Urbano",
    "condominio": "Urbano",
    "rancheria": "Rural",
    "rancho": "Rural",
    "ejido": "Rural",
    "congregacion": "Rural",
    "pueblo": "Semiurbano",
    "barrio": "Semiurbano",
}
ZONAS_VALIDAS = {"urbano": "Urbano", "semiurbano": "Semiurbano", "rural": "Rural"}


# --------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------


def normalizar_texto(valor) -> str:
    """Minúsculas, sin acentos, sin puntuación y con espacios simples."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[^0-9a-zA-Z]+", " ", texto.lower())
    return re.sub(r"\s+", " ", texto).strip()


def es_vacio(valor) -> bool:
    if valor is None:
        return True
    try:
        if pd.isna(valor):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(valor, str) and valor.strip() == ""


def clave(valor) -> str | None:
    """Normaliza un valor para usarlo como llave de cruce (12, 12.0, '12' → '12')."""
    if es_vacio(valor):
        return None
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, (int,)):
        return str(valor)
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else str(valor)
    texto = re.sub(r"\s+", " ", str(valor)).strip().upper()
    if re.fullmatch(r"-?\d+\.0+", texto):
        texto = texto.split(".")[0]
    return texto or None


def valor_limpio(valor):
    """Convierte floats enteros en int y NaN en None para la salida."""
    if es_vacio(valor):
        return None
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()
    return valor


def texto_llave(valor) -> str:
    """Representación en texto para concatenados (7601974.0 → '7601974')."""
    v = valor_limpio(valor)
    return "" if v is None else str(v).strip()


def inferir_campania(nombre_archivo: str | None) -> int | None:
    """'Cartera_Campaña_19.xlsx' → 19."""
    if not nombre_archivo:
        return None
    base = normalizar_texto(re.sub(r"\.[^.]+$", "", nombre_archivo))
    m = re.search(r"campana\s*(\d+)", base)
    if m:
        return int(m.group(1))
    numeros = re.findall(r"\d+", base)
    return int(numeros[-1]) if numeros else None


# --------------------------------------------------------------------------
# 4.1 Código postal
# --------------------------------------------------------------------------


def extraer_cp(direccion) -> str | None:
    if es_vacio(direccion):
        return None
    m = RE_CP.search(str(direccion).strip())
    if not m:
        return None
    cp = m.group(1).zfill(5)
    return None if cp == "00000" else cp


# --------------------------------------------------------------------------
# 4.3 / 4.4 Segmentación de la dirección
# --------------------------------------------------------------------------


@dataclass
class DireccionSegmentada:
    vialidad: str = ""
    no: str = ""
    interior: str = ""
    mza: str = ""
    lt: str = ""
    colonia: str = ""

    def direccion_calle(self) -> str:
        """Dirección calle en MAYÚSCULAS, ej. 'TULE NO. 79 INTERIOR B MZ 3 L- 12'."""
        partes = []
        if self.vialidad:
            partes.append(self.vialidad)
        if self.no:
            partes.append(f"No. {self.no}")
        if self.interior:
            partes.append(f"Interior {self.interior}")
        if self.mza:
            partes.append(f"Mz {self.mza}")
        if self.lt:
            partes.append(f"L- {self.lt}")
        return re.sub(r"\s+", " ", " ".join(partes)).upper()


def segmentar_direccion(direccion) -> DireccionSegmentada:
    """Separa Vialidad, No, Int, Mza, Lt y Colonia (mejor esfuerzo, sin inventar)."""
    seg = DireccionSegmentada()
    if es_vacio(direccion):
        return seg
    texto = RE_CP_FINAL.sub("", str(direccion)).strip()

    primera = RE_CLAVE.search(texto)
    if not primera:
        # Sin palabras clave: no se puede distinguir vialidad de colonia.
        seg.vialidad = re.sub(r"\s+", " ", texto).strip()
        return seg

    seg.vialidad = re.sub(r"\s+", " ", texto[: primera.start()]).strip()
    resto = texto[primera.start():]

    atributos = {"no": "no", "int": "interior", "mza": "mza", "lt": "lt"}
    for palabra in PALABRAS_CLAVE:
        m = re.search(rf"\b{palabra}\b\.?", resto, re.IGNORECASE)
        if not m:
            continue
        token = RE_TOKEN.match(resto, m.end())
        if token:
            setattr(seg, atributos[palabra.lower()], token.group(1))
            resto = resto[: m.start()] + " " + resto[token.end():]
        else:
            resto = resto[: m.start()] + " " + resto[m.end():]

    seg.colonia = re.sub(r"\s+", " ", resto).strip(" ,.-")
    return seg


# --------------------------------------------------------------------------
# Catálogo de códigos postales (SEPOMEX)
# --------------------------------------------------------------------------


def _buscar_columna(columnas, candidatos) -> str | None:
    norm = {normalizar_texto(c): c for c in columnas}
    for cand in candidatos:
        if cand in norm:
            return norm[cand]
    for cand in candidatos:
        for n, original in norm.items():
            if cand in n:
                return original
    return None


def zona_desde_tipo_asentamiento(tipo) -> str | None:
    t = normalizar_texto(tipo)
    if not t:
        return None
    if t in ZONAS_VALIDAS:
        return ZONAS_VALIDAS[t]
    for patron, zona in REGLAS_TIPO_ASENTAMIENTO.items():
        if t == patron or t.startswith(patron + " "):
            return zona
    return None


def _leer_txt_sepomex(contenido: bytes) -> pd.DataFrame:
    texto = None
    for codificacion in ("utf-8", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    lineas = texto.splitlines()
    inicio = 0
    for i, linea in enumerate(lineas[:20]):
        separadores = max(linea.count("|"), linea.count("\t"), linea.count(","))
        encabezado = normalizar_texto(linea)
        if separadores >= 2 and ("codigo" in encabezado or re.search(r"\bcp\b", encabezado)):
            inicio = i
            break
    separador = "|" if "|" in lineas[inicio] else ("\t" if "\t" in lineas[inicio] else ",")
    return pd.read_csv(
        io.StringIO("\n".join(lineas[inicio:])),
        sep=separador,
        dtype=str,
        keep_default_na=False,
        on_bad_lines="skip",
    )


def leer_catalogo_cp(contenido: bytes, nombre: str) -> pd.DataFrame:
    """Lee el catálogo SEPOMEX (TXT, XLS/XLSX con una hoja por estado, CSV o ZIP)
    y lo reduce a una fila por código postal: Cp, Municipio, Estado, Zona."""
    nombre_l = nombre.lower()
    if nombre_l.endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            internos = [n for n in z.namelist() if re.search(r"\.(txt|csv|xlsx?|)$", n.lower())]
            internos = [n for n in internos if not n.endswith("/")]
            if not internos:
                raise ValueError("El ZIP no contiene un archivo TXT/CSV/XLS del catálogo.")
            return leer_catalogo_cp(z.read(internos[0]), internos[0])

    if nombre_l.endswith((".xlsx", ".xls", ".xlsm")):
        hojas = pd.read_excel(io.BytesIO(contenido), sheet_name=None, dtype=str)
        partes = [h for n, h in hojas.items() if _buscar_columna(h.columns, ["d codigo", "codigo postal", "cp"])]
        if not partes:
            raise ValueError("No se encontró una columna de código postal en el catálogo.")
        crudo = pd.concat(partes, ignore_index=True)
    else:
        crudo = _leer_txt_sepomex(contenido)

    col_cp = _buscar_columna(crudo.columns, ["d codigo", "codigo postal", "cp", "codigo"])
    col_mun = _buscar_columna(crudo.columns, ["d mnpio", "municipio", "mnpio", "poblacion"])
    col_edo = _buscar_columna(crudo.columns, ["d estado", "estado"])
    col_zona = _buscar_columna(crudo.columns, ["d zona", "zona"])
    col_tipo = _buscar_columna(crudo.columns, ["d tipo asenta", "tipo asentamiento", "tipo asenta"])
    if not col_cp:
        raise ValueError("No se encontró la columna de código postal en el catálogo.")

    df = pd.DataFrame({"Cp": crudo[col_cp].astype(str).str.strip()})
    df = df.assign(
        Municipio=crudo[col_mun].astype(str).str.strip() if col_mun else None,
        Estado=crudo[col_edo].astype(str).str.strip() if col_edo else None,
        _zona=crudo[col_zona] if col_zona else None,
        _tipo=crudo[col_tipo] if col_tipo else None,
    )
    df = df[df["Cp"].str.fullmatch(r"\d{1,5}", na=False)].copy()
    df["Cp"] = df["Cp"].str.zfill(5)

    def zona_fila(fila):
        z = ZONAS_VALIDAS.get(normalizar_texto(fila["_zona"]))
        return z or zona_desde_tipo_asentamiento(fila["_tipo"])

    df["Zona"] = df.apply(zona_fila, axis=1) if len(df) else pd.Series(dtype=object)

    def primero(serie):
        s = serie.dropna()
        s = s[s.astype(str).str.strip() != ""]
        return s.mode().iloc[0] if len(s) else None

    catalogo = (
        df.groupby("Cp", sort=True)
        .agg(Municipio=("Municipio", primero), Estado=("Estado", primero), Zona=("Zona", primero))
        .reset_index()
    )
    return catalogo


# --------------------------------------------------------------------------
# Estructura General de Bases
# --------------------------------------------------------------------------


@dataclass
class EstructuraGeneral:
    zonas: pd.DataFrame  # ZONA, REGION, DIVISION, RUTA, ID COBRADOR, _clave
    calendario: pd.DataFrame  # hoja cruda (header=None)
    campanias: pd.DataFrame  # hoja cruda (header=None)
    advertencias: list[str] = field(default_factory=list)

    def campanias_disponibles(self) -> list[int]:
        a = set(_bloques(self.calendario))
        b = set(_bloques(self.campanias))
        return sorted(a & b) if a and b else sorted(a | b)


def _hoja(hojas: dict, nombre: str) -> pd.DataFrame:
    objetivo = normalizar_texto(nombre)
    for n, df in hojas.items():
        if normalizar_texto(n) == objetivo:
            return df
    for n, df in hojas.items():
        if objetivo in normalizar_texto(n):
            return df
    raise ValueError(f"No se encontró la hoja '{nombre}' en el archivo de Estructura General.")


def _bloques(crudo: pd.DataFrame, filas_busqueda: int = 10) -> dict[int, tuple[int, int]]:
    """Encuentra encabezados 'Campaña de Trabajo N' → {N: (fila, columna)}."""
    encontrados: dict[int, tuple[int, int]] = {}
    for r in range(min(filas_busqueda, len(crudo))):
        for c in range(crudo.shape[1]):
            m = re.fullmatch(r"campana de trabajo\s*(\d+)", normalizar_texto(crudo.iat[r, c]))
            if m and int(m.group(1)) not in encontrados:
                encontrados[int(m.group(1))] = (r, c)
    return encontrados


def _rango_bloque(crudo: pd.DataFrame, n: int, hoja: str) -> tuple[int, int, int]:
    bloques = _bloques(crudo)
    if n not in bloques:
        disponibles = ", ".join(str(k) for k in sorted(bloques)) or "ninguna"
        raise ValueError(
            f"No existe el bloque 'Campaña de Trabajo {n}' en la hoja '{hoja}'. "
            f"Campañas disponibles: {disponibles}."
        )
    fila, col = bloques[n]
    siguientes = [c for (r, c) in bloques.values() if r == fila and c > col]
    fin = min(siguientes) if siguientes else crudo.shape[1]
    return fila, col, fin


def _fila_subencabezado(crudo, fila, col, fin, palabras) -> int:
    for r in range(fila + 1, min(fila + 4, len(crudo))):
        textos = [normalizar_texto(crudo.iat[r, c]) for c in range(col, fin)]
        if any(any(p in t for p in palabras) for t in textos):
            return r
    return fila + 1


def tabla_calendario(estructura: EstructuraGeneral, n: int) -> dict[str, object]:
    """RUTA → Fecha Cierre para la Campaña de Trabajo N (sección 4.7)."""
    crudo = estructura.calendario
    fila, col, fin = _rango_bloque(crudo, n, "Calendario de Cierre")
    sub = _fila_subencabezado(crudo, fila, col, fin, ["ruta", "fecha"])
    nombres = {c: normalizar_texto(crudo.iat[sub, c]) for c in range(col, fin)}

    col_ruta = next((c for c, t in nombres.items() if t == "ruta"), None)
    col_ruta = col_ruta if col_ruta is not None else next((c for c, t in nombres.items() if "ruta" in t), col)
    col_fecha = next((c for c, t in nombres.items() if "cierre" in t), None)
    if col_fecha is None:
        fechas = [c for c, t in nombres.items() if "fecha" in t]
        col_fecha = fechas[-1] if fechas else min(col + 2, fin - 1)

    tabla: dict[str, object] = {}
    for r in range(sub + 1, len(crudo)):
        k = clave(crudo.iat[r, col_ruta])
        if k is None or k in tabla:
            continue
        tabla[k] = valor_limpio(crudo.iat[r, col_fecha])
    return tabla


def tabla_morosidad(estructura: EstructuraGeneral, n: int) -> dict[str, object]:
    """CampaniaSaldo → Mora para la Campaña de Trabajo N (sección 4.8)."""
    crudo = estructura.campanias
    fila, col, fin = _rango_bloque(crudo, n, "Campaña de Trabajo")
    sub = _fila_subencabezado(crudo, fila, col, fin, ["campana", "mora"])
    nombres = {c: normalizar_texto(crudo.iat[sub, c]) for c in range(col, fin)}

    col_camp = next((c for c, t in nombres.items() if "campana" in t), col)
    col_mora = next((c for c, t in nombres.items() if "mora" in t), min(col + 1, fin - 1))

    tabla: dict[str, object] = {}
    for r in range(sub + 1, len(crudo)):
        k = clave(crudo.iat[r, col_camp])
        if k is None or k in tabla:
            continue
        tabla[k] = valor_limpio(crudo.iat[r, col_mora])
    return tabla


def _leer_base_zonas(crudo: pd.DataFrame, advertencias: list[str]) -> pd.DataFrame:
    fila_enc = None
    for r in range(min(15, len(crudo))):
        if "zona" in [normalizar_texto(v) for v in crudo.iloc[r].tolist()]:
            fila_enc = r
            break
    if fila_enc is None:
        raise ValueError("La hoja 'Base de Zonas' no tiene una columna 'ZONA'.")

    encabezados = [str(v) if not es_vacio(v) else f"_col{i}" for i, v in enumerate(crudo.iloc[fila_enc])]
    df = crudo.iloc[fila_enc + 1:].copy()
    df.columns = encabezados

    def col(cands):
        c = _buscar_columna(df.columns, cands)
        if c is None:
            raise ValueError(f"En 'Base de Zonas' falta la columna {cands[0].upper()}.")
        return c

    zonas = pd.DataFrame(
        {
            "ZONA": df[col(["zona"])],
            "REGION": df[col(["region"])],
            "DIVISION": df[col(["division"])],
            "RUTA": df[col(["ruta"])],
            "ID COBRADOR": df[col(["no cobrador", "id cobrador", "cobrador"])],
        }
    )
    zonas["_clave"] = zonas["ZONA"].map(clave)
    zonas = zonas[zonas["_clave"].notna()]
    duplicadas = zonas["_clave"][zonas["_clave"].duplicated()].unique()
    if len(duplicadas):
        advertencias.append(
            f"'Base de Zonas' tiene {len(duplicadas)} ZONA(s) duplicada(s); se usó la primera "
            f"aparición: {', '.join(map(str, duplicadas[:10]))}{'…' if len(duplicadas) > 10 else ''}"
        )
    return zonas.drop_duplicates("_clave", keep="first").reset_index(drop=True)


def leer_estructura(contenido: bytes) -> EstructuraGeneral:
    hojas = pd.read_excel(io.BytesIO(contenido), sheet_name=None, header=None)
    advertencias: list[str] = []
    zonas = _leer_base_zonas(_hoja(hojas, "Base de Zonas"), advertencias)
    return EstructuraGeneral(
        zonas=zonas,
        calendario=_hoja(hojas, "Calendario de Cierre"),
        campanias=_hoja(hojas, "Campaña de Trabajo"),
        advertencias=advertencias,
    )


# --------------------------------------------------------------------------
# Cartera
# --------------------------------------------------------------------------


def elegir_hoja_cartera(nombres: list[str]) -> str:
    """Hoja 'BASE' si existe; si no, la primera."""
    return next((h for h in nombres if normalizar_texto(h) == "base"), nombres[0])


def ubicar_encabezado(filas) -> int:
    """Índice (0-based) de la fila de encabezados: la primera que contiene 'NoDama'."""
    for i, fila in enumerate(filas):
        if i >= 20:
            break
        if any(normalizar_texto(v) == "nodama" for v in fila):
            return i
    return 0


def es_excel_openpyxl(nombre: str) -> bool:
    return nombre.lower().endswith((".xlsx", ".xlsm"))


def leer_cartera(contenido: bytes, nombre: str = "") -> pd.DataFrame:
    """Lee la cartera. La columna '_fila_excel' guarda la fila original de cada cuenta en
    la hoja, para poder devolver después el mismo archivo con las columnas anexadas."""
    if nombre.lower().endswith(".csv"):
        df = pd.read_csv(io.BytesIO(contenido), dtype=object)
        df["_fila_excel"] = None
        return _normalizar_columnas_cartera(df)
    if not es_excel_openpyxl(nombre):  # .xls
        hojas = pd.ExcelFile(io.BytesIO(contenido)).sheet_names
        df = pd.read_excel(io.BytesIO(contenido), sheet_name=elegir_hoja_cartera(hojas), dtype=object)
        df["_fila_excel"] = None
        return _normalizar_columnas_cartera(df)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        ws = wb[elegir_hoja_cartera(wb.sheetnames)]
        filas = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not filas:
        raise ValueError("La hoja de la cartera está vacía.")
    h = ubicar_encabezado(filas)
    encabezados, vistos = [], {}
    for i, v in enumerate(filas[h]):
        nombre_col = str(v).strip() if not es_vacio(v) else f"_col{i + 1}"
        if nombre_col in vistos:
            vistos[nombre_col] += 1
            nombre_col = f"{nombre_col}.{vistos[nombre_col]}"
        else:
            vistos[nombre_col] = 0
        encabezados.append(nombre_col)
    ancho = len(encabezados)
    datos = [list(f[:ancho]) + [None] * (ancho - len(f)) for f in filas[h + 1:]]
    df = pd.DataFrame(datos, columns=encabezados, dtype=object)
    df["_fila_excel"] = list(range(h + 2, h + 2 + len(datos)))
    return _normalizar_columnas_cartera(df)


def _normalizar_columnas_cartera(df: pd.DataFrame) -> pd.DataFrame:
    """Renombra columnas a su nombre canónico sin importar mayúsculas/acentos.
    Si existe 'ZONA' y un duplicado 'Zona', se conserva 'ZONA'."""
    df = df.copy()
    renombres = {}
    usados = set()
    for canon in COLUMNAS_CARTERA:
        objetivo = normalizar_texto(canon)
        candidatas = [c for c in df.columns if normalizar_texto(c) == objetivo and c not in usados]
        if not candidatas:
            continue
        # Preferir coincidencia exacta (p.ej. 'ZONA' sobre 'Zona')
        elegida = canon if canon in candidatas else candidatas[0]
        renombres[elegida] = canon
        usados.update(candidatas)
    df = df.rename(columns=renombres)
    faltantes = [c for c in COLUMNAS_CARTERA if c not in df.columns]
    if faltantes:
        raise ValueError("La cartera no contiene las columnas: " + ", ".join(faltantes))
    df = df[~df[COLUMNAS_CARTERA].map(es_vacio).all(axis=1)]
    return df.reset_index(drop=True)


# --------------------------------------------------------------------------
# Proceso principal
# --------------------------------------------------------------------------


@dataclass
class Resultado:
    base: pd.DataFrame  # COLUMNAS_BASE
    resumen: dict
    advertencias: list[str]


def generar_base(
    cartera: pd.DataFrame,
    estructura: EstructuraGeneral,
    catalogo_cp: pd.DataFrame | None,
    campania_trabajo: int,
) -> Resultado:
    advertencias = list(estructura.advertencias)
    calendario = tabla_calendario(estructura, campania_trabajo)
    morosidad = tabla_morosidad(estructura, campania_trabajo)
    zonas = estructura.zonas.set_index("_clave")[["REGION", "RUTA", "DIVISION", "ID COBRADOR"]].to_dict("index")

    cps: dict[str, dict] = {}
    if catalogo_cp is not None and len(catalogo_cp):
        cps = catalogo_cp.set_index("Cp")[["Municipio", "Estado", "Zona"]].to_dict("index")
    else:
        advertencias.append(
            "No hay catálogo de códigos postales cargado: Municipio, Estado y Zona (Urbano/Rural) quedarán vacíos."
        )

    filas = []
    for _, r in cartera.iterrows():
        motivos = []

        # 4.1 / 4.2
        cp = extraer_cp(r["Direccion"])
        info_cp = cps.get(cp) if cp else None
        if not cp:
            motivos.append("CP no identificado")
        elif not info_cp:
            motivos.append("CP no encontrado en catálogo")

        # 4.3 / 4.4
        seg = segmentar_direccion(r["Direccion"])

        # 4.5
        info_zona = zonas.get(clave(r["ZONA"]))
        if not info_zona:
            motivos.append("ZONA no encontrada en Base de Zonas")
        ruta = valor_limpio(info_zona["RUTA"]) if info_zona else None

        # 4.7
        fecha_cierre = calendario.get(clave(ruta)) if ruta is not None else None
        if fecha_cierre is None:
            motivos.append(f"Ruta sin Fecha de Cierre en Campaña {campania_trabajo}")

        # 4.8
        mora = morosidad.get(clave(r["CampaniaSaldo"]))
        if mora is None:
            motivos.append(f"CampaniaSaldo sin Morosidad en Campaña {campania_trabajo}")

        nodama = texto_llave(r["NoDama"])
        filas.append(
            {
                "ZONA": valor_limpio(r["ZONA"]),
                "REGION": valor_limpio(info_zona["REGION"]) if info_zona else None,
                "RUTA": ruta,
                "DIVISION": valor_limpio(info_zona["DIVISION"]) if info_zona else None,
                "ID COBRADOR": valor_limpio(info_zona["ID COBRADOR"]) if info_zona else None,
                "NoDama": valor_limpio(r["NoDama"]),
                "Direccion": valor_limpio(r["Direccion"]),
                "Direccion Calle": seg.direccion_calle() or None,
                "Colonia": seg.colonia or None,
                "Municipio / Poblacion": info_cp["Municipio"] if info_cp else None,
                "Cp": cp,
                "Estado": info_cp["Estado"] if info_cp else None,
                "Zona (Urbano/Rural)": info_cp["Zona"] if info_cp else None,
                "Referencia": valor_limpio(r["Referencia"]),
                "AnioSaldo": valor_limpio(r["AnioSaldo"]),
                "CampaniaSaldo": valor_limpio(r["CampaniaSaldo"]),
                "DigitoVerificador": valor_limpio(r["DigitoVerificador"]),
                "Concatenado": f"{nodama}-{texto_llave(r['CampaniaSaldo'])}",
                "Fecha de cierre": fecha_cierre,
                "Morosidad": mora,
                "Campaña de trabajo": campania_trabajo,
                "Referencia de Pago": f"{nodama}-{texto_llave(r['DigitoVerificador'])}",
                "requiere_revision": bool(motivos),
                "motivo_revision": "; ".join(motivos),
                "_fila_excel": valor_limpio(r["_fila_excel"]) if "_fila_excel" in r else None,
            }
        )

    base = pd.DataFrame(filas, columns=COLUMNAS_BASE, dtype=object)
    base["requiere_revision"] = base["requiere_revision"].astype(bool)
    total = len(base)
    cp_ok = int(base["Cp"].notna().sum())
    motivos = base["motivo_revision"].astype(str)
    cp_cat = int((~motivos.str.contains("CP no", regex=False)).sum())
    zona_ok = int((~motivos.str.contains("ZONA no encontrada", regex=False)).sum())
    fecha_ok = int(base["Fecha de cierre"].notna().sum())
    mora_ok = int(base["Morosidad"].notna().sum())
    resumen = {
        "Total de filas procesadas": total,
        "CP identificado": cp_ok,
        "CP no identificado": total - cp_ok,
        "CP encontrado en catálogo": cp_cat,
        "CP no encontrado en catálogo": cp_ok - cp_cat,
        "Zona encontrada": zona_ok,
        "Zona no encontrada": total - zona_ok,
        "Fecha de cierre calculada": fecha_ok,
        "Fecha de cierre no calculada": total - fecha_ok,
        "Morosidad calculada": mora_ok,
        "Morosidad no calculada": total - mora_ok,
        "Filas que requieren revisión": int(base["requiere_revision"].sum()),
    }
    return Resultado(base=base, resumen=resumen, advertencias=advertencias)


# --------------------------------------------------------------------------
# Exportación
# --------------------------------------------------------------------------


def exportar_excel(resultado: Resultado, campania_trabajo: int) -> bytes:
    """Excel con la base (filas a revisar en amarillo), hoja de revisión y resumen."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    base = resultado.base
    salida = base[COLUMNAS_SALIDA]
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        salida.to_excel(writer, sheet_name="Base de Gestion", index=False)
        revision = base[base["requiere_revision"]][["NoDama", "ZONA", "Direccion", "Cp", "motivo_revision"]]
        revision = revision.rename(columns={"motivo_revision": "Motivo de revisión"})
        revision.to_excel(writer, sheet_name="Revision", index=False)
        resumen = pd.DataFrame(
            [("Campaña de trabajo", campania_trabajo), ("Fecha de generación", datetime.now().strftime("%d/%m/%Y %H:%M"))]
            + list(resultado.resumen.items()),
            columns=["Concepto", "Valor"],
        )
        resumen.to_excel(writer, sheet_name="Resumen", index=False)

        ws = writer.sheets["Base de Gestion"]
        amarillo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ncols = len(COLUMNAS_SALIDA)
        for i, revisar in enumerate(base["requiere_revision"].tolist(), start=2):
            if revisar:
                for c in range(1, ncols + 1):
                    ws.cell(row=i, column=c).fill = amarillo
        col_fecha = COLUMNAS_SALIDA.index("Fecha de cierre") + 1
        for fila in ws.iter_rows(min_row=2, min_col=col_fecha, max_col=col_fecha):
            for celda in fila:
                celda.number_format = "DD/MM/YYYY"
        for c in range(1, ncols + 1):
            ws.cell(row=1, column=c).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(c)].width = 18
        for nombre in ("Direccion", "Direccion Calle", "Colonia", "Referencia"):
            letra = get_column_letter(COLUMNAS_SALIDA.index(nombre) + 1)
            ws.column_dimensions[letra].width = 40
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for nombre_hoja, anchos in (("Revision", [14, 10, 60, 8, 70]), ("Resumen", [34, 22])):
            h = writer.sheets[nombre_hoja]
            for i, ancho in enumerate(anchos, start=1):
                h.column_dimensions[get_column_letter(i)].width = ancho
                h.cell(row=1, column=i).font = Font(bold=True)
    return buffer.getvalue()


def exportar_csv(resultado: Resultado) -> bytes:
    df = resultado.base.drop(columns=["_fila_excel"], errors="ignore")
    df["Fecha de cierre"] = pd.to_datetime(df["Fecha de cierre"], errors="coerce").dt.strftime("%d/%m/%Y")
    return df.to_csv(index=False).encode("utf-8-sig")


# --------------------------------------------------------------------------
# Exportación sobre el archivo original de la cartera
# --------------------------------------------------------------------------

# Columnas que calcula el sistema (las demás de COLUMNAS_SALIDA vienen de la cartera)
COLUMNAS_SISTEMA = [c for c in COLUMNAS_SALIDA if c not in COLUMNAS_CARTERA]
COLUMNA_MOTIVO = "Motivo de revisión"
COLOR_SISTEMA = "#BDD7EE"  # azul claro por defecto
COLOR_REVISION = "#FFFF00"


def _argb(color: str) -> str:
    c = color.strip().lstrip("#").upper()
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    if not re.fullmatch(r"[0-9A-F]{6}", c):
        raise ValueError(f"Color inválido: {color}")
    return "FF" + c


def _nombre_hoja_libre(wb, nombre: str) -> str:
    candidato, i = nombre, 2
    while candidato in wb.sheetnames:
        candidato, i = f"{nombre} ({i})", i + 1
    return candidato


def exportar_excel_original(
    contenido: bytes,
    nombre: str,
    resultado: Resultado,
    campania_trabajo: int,
    color_sistema: str = COLOR_SISTEMA,
    colorear_llenadas: bool = False,
) -> bytes:
    """Devuelve la cartera tal como se envió (mismas columnas, orden, colores y formato),
    llenando las columnas vacías que el sistema calcula y anexando al final las que no
    existían. Las columnas anexadas llevan ``color_sistema``; las celdas anexadas de las
    filas que requieren revisión van en amarillo. Si ``colorear_llenadas`` es True, las
    columnas que ya existían y llenó el sistema también se pintan con ``color_sistema``."""
    from copy import copy

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not es_excel_openpyxl(nombre):
        raise ValueError("Sólo se puede conservar el formato de archivos .xlsx o .xlsm.")

    base = resultado.base
    if base["_fila_excel"].isna().any():
        raise ValueError("La base no tiene la ubicación de las filas en el archivo original.")

    wb = load_workbook(io.BytesIO(contenido), keep_vba=nombre.lower().endswith(".xlsm"))
    ws = wb[elegir_hoja_cartera(wb.sheetnames)]
    filas_enc = [
        [c.value for c in fila]
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20))
    ]
    fila_enc = ubicar_encabezado(filas_enc) + 1  # 1-based

    encabezados = {}
    ultima = 0
    for celda in ws[fila_enc]:
        if not es_vacio(celda.value):
            ultima = max(ultima, celda.column)
            encabezados.setdefault(normalizar_texto(celda.value), celda.column)

    relleno_sistema = PatternFill(start_color=_argb(color_sistema), end_color=_argb(color_sistema), fill_type="solid")
    relleno_revision = PatternFill(start_color=_argb(COLOR_REVISION), end_color=_argb(COLOR_REVISION), fill_type="solid")
    modelo_enc = ws.cell(row=fila_enc, column=max(ultima, 1))

    destino: dict[str, int] = {}
    anexadas: set[int] = set()
    siguiente = ultima + 1
    for col in COLUMNAS_SISTEMA + [COLUMNA_MOTIVO]:
        existente = encabezados.get(normalizar_texto(col))
        if existente and col != COLUMNA_MOTIVO:
            destino[col] = existente
            continue
        destino[col] = siguiente
        anexadas.add(siguiente)
        enc = ws.cell(row=fila_enc, column=siguiente, value=col)
        enc.font = copy(modelo_enc.font) if modelo_enc.has_style else Font(bold=True)
        enc.border = copy(modelo_enc.border)
        enc.alignment = copy(modelo_enc.alignment)
        enc.fill = relleno_sistema
        ancho = 40 if col in ("Direccion Calle", "Colonia", COLUMNA_MOTIVO) else 18
        ws.column_dimensions[get_column_letter(siguiente)].width = ancho
        siguiente += 1

    pintar = set(anexadas)
    if colorear_llenadas:
        pintar |= {c for n, c in destino.items() if c not in anexadas}
        for c in pintar - anexadas:
            ws.cell(row=fila_enc, column=c).fill = relleno_sistema

    col_fecha = destino["Fecha de cierre"]
    for fila in base.to_dict("records"):
        r = int(fila["_fila_excel"])
        revisar = bool(fila["requiere_revision"])
        for col, c in destino.items():
            valor = (fila["motivo_revision"] or None) if col == COLUMNA_MOTIVO else fila[col]
            celda = ws.cell(row=r, column=c, value=valor)
            if c == col_fecha and valor is not None:
                celda.number_format = "DD/MM/YYYY"
            if c in anexadas:
                celda.fill = relleno_revision if revisar else relleno_sistema
            elif c in pintar:
                celda.fill = relleno_sistema

    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{fila_enc}:{get_column_letter(siguiente - 1)}{ws.max_row}"

    # Hojas adicionales de revisión y resumen (las hojas originales no se tocan)
    rev = wb.create_sheet(_nombre_hoja_libre(wb, "Revision"))
    rev.append(["Fila", "NoDama", "ZONA", "Direccion", "Cp", COLUMNA_MOTIVO])
    for fila in base[base["requiere_revision"]].to_dict("records"):
        rev.append([fila["_fila_excel"], fila["NoDama"], fila["ZONA"], fila["Direccion"], fila["Cp"], fila["motivo_revision"]])
    res = wb.create_sheet(_nombre_hoja_libre(wb, "Resumen"))
    res.append(["Concepto", "Valor"])
    res.append(["Campaña de trabajo", campania_trabajo])
    res.append(["Fecha de generación", datetime.now().strftime("%d/%m/%Y %H:%M")])
    for k, v in resultado.resumen.items():
        res.append([k, v])
    for hoja, anchos in ((rev, [8, 14, 10, 60, 8, 70]), (res, [34, 22])):
        for i, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(i)].width = ancho
            hoja.cell(row=1, column=i).font = Font(bold=True)

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# --------------------------------------------------------------------------
# Base para visitas de gestores
# --------------------------------------------------------------------------

PLANTILLA_VISITAS = Path(__file__).parent / "plantillas" / "Base_para_visitas.xlsx"

COLUMNAS_VISITAS = [
    "FECHA DE ASIGNACION",
    "ZONA",
    "ASIGNACION",  # única columna que se entrega en blanco (la llena el supervisor)
    "NoDama",
    "DIGITO VERIFICADOR",
    "NOMBRE",
    "DIRECCION",
    "COLONIA",
    "CP Extraido",
    "LOCALIDAD",
    "REFERENCIA",
    "TEMPORALIDAD",
    "CAMPANA",
    "IMPORTE NETO FACTURA",
    "TELEFONO CELULAR",
    "DescSituacionCie",
]

# Columnas de visitas que se toman directo de la cartera: nombres aceptados en orden de
# preferencia (se comparan sin acentos, mayúsculas, espacios ni signos: "SaldoDama" = "Saldo Dama").
CAMPOS_CARTERA_VISITAS = {
    "NOMBRE": ["nombre", "nombrecliente", "nombrecompleto", "nombredama", "cliente"],
    "IMPORTE NETO FACTURA": ["saldodama", "importenetofactura", "importeneto", "importefactura"],
    "TELEFONO CELULAR": ["telefonocelular", "celular", "telcelular"],
    "DescSituacionCie": ["descsituacioncie", "situacioncie", "descsituacion"],
}
# Nombre de la columna de la cartera que se muestra en los avisos
FUENTE_VISITAS = {
    "NOMBRE": "NOMBRE",
    "IMPORTE NETO FACTURA": "SaldoDama",
    "TELEFONO CELULAR": "TelefonoCelular",
    "DescSituacionCie": "DescSituacionCie",
}

MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
RE_PREFIJO_COLONIA = re.compile(r"^(COLONIA|COL)\b\.?\s*", re.IGNORECASE)


def _compacto(valor) -> str:
    return normalizar_texto(valor).replace(" ", "")


def _columna_cartera(columnas, candidatos) -> str | None:
    norm = {}
    for c in columnas:
        if not str(c).startswith("_"):
            norm.setdefault(_compacto(c), c)
    for cand in candidatos:
        if cand in norm:
            return norm[cand]
    for n, original in norm.items():
        if n.startswith(candidatos[0]):
            return original
    return None


def columnas_visitas_faltantes(cartera: pd.DataFrame | None) -> list[str]:
    """Columnas de visitas que dependen de la cartera y no vienen en ella."""
    if cartera is None:
        return list(CAMPOS_CARTERA_VISITAS)
    return [k for k, cands in CAMPOS_CARTERA_VISITAS.items() if _columna_cartera(cartera.columns, cands) is None]


def _colonia_visitas(colonia) -> str | None:
    if es_vacio(colonia):
        return None
    return RE_PREFIJO_COLONIA.sub("", str(colonia)).strip() or None


def _localidad(municipio) -> str | None:
    if es_vacio(municipio):
        return None
    texto = unicodedata.normalize("NFKD", str(municipio))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).replace(".", "")
    return re.sub(r"\s+", " ", texto).strip().upper()


def _temporalidad(mora) -> str | None:
    m = valor_limpio(mora)
    if m is None:
        return None
    return str(m) if normalizar_texto(m).startswith("mora") else f"Mora {m}"


def _campana(anio, campania):
    """AnioSaldo 2025 + CampaniaSaldo 12 → 202512."""
    a, c = clave(anio), clave(campania)
    if a and c and a.isdigit() and c.isdigit():
        return int(f"{a}{int(c):02d}")
    return None


def _cp_numero(cp):
    return int(cp) if cp and str(cp).isdigit() else None


def tabla_visitas(
    resultado: Resultado, cartera: pd.DataFrame | None, fecha_asignacion
) -> pd.DataFrame:
    """Arma la base para visitas: una fila por cuenta, todas las columnas llenas con datos de
    la base (y de la cartera original para nombre, importe, teléfono y situación), excepto
    ASIGNACION, que se entrega en blanco."""
    base = resultado.base.reset_index(drop=True)
    extras = pd.DataFrame(index=base.index)
    if cartera is not None:
        cart = cartera.reset_index(drop=True)
        if (
            "_fila_excel" in cart
            and cart["_fila_excel"].notna().all()
            and base["_fila_excel"].notna().all()
        ):
            cart = cart.set_index("_fila_excel").reindex(base["_fila_excel"].astype(int)).reset_index(drop=True)
        elif len(cart) != len(base):
            cart = None
        if cart is not None:
            for campo, cands in CAMPOS_CARTERA_VISITAS.items():
                col = _columna_cartera(cart.columns, cands)
                if col is not None:
                    extras[campo] = cart[col].map(valor_limpio)

    fecha = pd.Timestamp(fecha_asignacion).to_pydatetime() if fecha_asignacion is not None else None
    filas = []
    registros_extra = extras.to_dict("records") if len(extras.columns) else [{}] * len(base)
    for f, extra in zip(base.to_dict("records"), registros_extra):
        filas.append(
            {
                "FECHA DE ASIGNACION": fecha,
                "ZONA": f["ZONA"],
                "ASIGNACION": None,
                "NoDama": f["NoDama"],
                "DIGITO VERIFICADOR": f["DigitoVerificador"],
                "NOMBRE": extra.get("NOMBRE"),
                # Dirección que genera el sistema (Direccion Calle); si no se pudo armar, la original
                "DIRECCION": f["Direccion Calle"] or f["Direccion"],
                "COLONIA": _colonia_visitas(f["Colonia"]),
                "CP Extraido": _cp_numero(f["Cp"]),
                "LOCALIDAD": _localidad(f["Municipio / Poblacion"]),
                "REFERENCIA": f["Referencia"],
                "TEMPORALIDAD": _temporalidad(f["Morosidad"]),
                "CAMPANA": _campana(f["AnioSaldo"], f["CampaniaSaldo"]),
                "IMPORTE NETO FACTURA": extra.get("IMPORTE NETO FACTURA"),
                "TELEFONO CELULAR": extra.get("TELEFONO CELULAR"),
                "DescSituacionCie": extra.get("DescSituacionCie"),
            }
        )
    return pd.DataFrame(filas, columns=COLUMNAS_VISITAS, dtype=object)


def nombre_hoja_visitas(fecha_asignacion) -> str:
    if fecha_asignacion is None:
        return "VISITAS"
    f = pd.Timestamp(fecha_asignacion)
    return f"VISITAS {f.day} {MESES[f.month - 1]}"


def exportar_visitas(tabla: pd.DataFrame, fecha_asignacion=None) -> bytes:
    """Escribe la base de visitas sobre la plantilla (mismo encabezado, colores, anchos y
    formato de fecha que la base de visitas que usan los gestores)."""
    from copy import copy

    from openpyxl import load_workbook

    wb = load_workbook(PLANTILLA_VISITAS)
    ws = wb.active
    ws.title = nombre_hoja_visitas(fecha_asignacion)
    encabezados = [c.value for c in ws[1]]
    if encabezados[: len(COLUMNAS_VISITAS)] != COLUMNAS_VISITAS:
        raise ValueError("La plantilla de visitas no coincide con las columnas esperadas.")
    # Estilo de la fila modelo (fila 2 de la plantilla); copiar el índice de estilo es mucho más
    # rápido que copiar fuente, relleno, bordes, etc. celda por celda.
    estilos = {c.column: c._style for c in ws[2]}
    for i, fila in enumerate(tabla[COLUMNAS_VISITAS].itertuples(index=False), start=2):
        for j, valor in enumerate(fila, start=1):
            celda = ws.cell(row=i, column=j, value=valor_limpio(valor))
            celda._style = copy(estilos[j])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
