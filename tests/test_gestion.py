import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook

import procesamiento as proc
from tests.test_procesamiento import _estructura_xlsx

CATALOGO = pd.DataFrame(
    {"Cp": ["45679"], "Municipio": ["Tlajomulco de Zúñiga"], "Estado": ["Jalisco"], "Zona": ["Rural"]}
)
ENC_CARTERA = [
    "Zona", "NoDama", "Nombre", "Direccion", "Referencia", "TelefonoCasa", "AnioSaldo", "CampaniaSaldo",
    "AnioCampaniaSaldo", "ImporteNetoFactura", "SaldoDama", "DigitoVerificador", "MotivoNoCobro", "Cancelacion",
    "PrimeraOrden", "Reactivacion", "TelefonoCelular", "IdSituacion", "DescSituacion", "IdSituacionCie",
    "DescSituacionCie", "TipoNombramiento", "Geolocalizacion", "NumeroLiquidacion", "id",
    "REGION", "RUTA", "DIVISION", "ID COBRADOR", "Concatenado", "Fecha de cierre", "Morosidad",
    "Campaña de trabajo", "Referencia de Pago", "ColumnaNueva",
]


def _cartera() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BASE"
    ws.append(ENC_CARTERA)
    ws.append([101, 7870740, "MA DEL ROCIO REYES GONZALEZ", "TULE No 79 COLONIA ALAMEDA CP 45679",
               "INDEPENDENCIA Y JUAREZ", 0, 2026, 15, "2026,15", 276, 276, 3, 17, None, None, None,
               3333381020, 1, "ENTREGADO", 0, None, None, None, 2, -37] + [None] * 9 + ["x"])
    ws.append([999, 7872443, "VICTORIA HERNANDEZ MENCHACA", "PUERTO ALTATA No 1 COLONIA EL MUELLE CP 45683",
               "20 NOVIEMBRE", 0, 2026, 17, "2026,17", 1275, 875, 2, 2, "C", None, None,
               3313401850, 2, "NO HUBO QUIEN RECIBIERA", 1, "ENTREGADO EN BODEGA", None, None, 2, -40] + [None] * 10)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generar():
    contenido = _cartera()
    cartera = proc.leer_cartera(contenido, "Cartera_Campaña_19.xlsx")
    res = proc.generar_base(cartera, proc.leer_estructura(_estructura_xlsx()), CATALOGO, 19)
    return load_workbook(io.BytesIO(proc.exportar_excel_gestion(res, cartera, 19)))


def test_formato_de_la_plantilla():
    wb = _generar()
    ws = wb["BASE"]
    enc = [c.value for c in ws[1]]
    plantilla = [c.value for c in load_workbook(proc.PLANTILLA_GESTION).active[1]]
    assert enc[:40] == plantilla and enc[40:] == ["ColumnaNueva"]
    # Colores de encabezado de la plantilla
    assert ws["A1"].fill.fgColor.theme == 8 and ws["G1"].fill.fgColor.rgb == "FF538ED5"
    assert ws["AM1"].fill.fgColor.theme == 5 and ws["AO1"].fill.fgColor.rgb == "FF538ED5"
    assert ws["A2"].font.name == "Aptos Narrow" and ws["AJ2"].number_format == "mm-dd-yy"
    assert wb.sheetnames == ["BASE", "Revision", "Resumen"]


def test_cada_dato_bajo_su_encabezado():
    ws = _generar()["BASE"]
    enc = [c.value.strip() for c in ws[1]]
    f = dict(zip(enc, [c.value for c in ws[2]]))
    assert f["ZONA"] == 101 and f["REGION"] == "CENTRO" and f["RUTA"] == 5 and f["ID COBRADOR"] == 9001
    assert f["Zona"] == 101 and f["NoDama"] == 7870740 and f["Nombre"] == "MA DEL ROCIO REYES GONZALEZ"
    assert f["Direccion Calle"] == "TULE NO. 79" and f["Colonia"] == "COLONIA ALAMEDA"
    assert f["Municipio / Poblacion"] == "Tlajomulco de Zúñiga" and f["Cp"] == "45679" and f["Estado"] == "Jalisco"
    assert f["TelefonoCasa"] == 0 and f["AnioSaldo"] == 2026 and f["CampaniaSaldo"] == 15
    assert f["AnioCampaniaSaldo"] == "2026,15" and f["ImporteNetoFactura"] == 276 and f["SaldoDama"] == 276
    assert f["DigitoVerificador"] == 3 and f["TelefonoCelular"] == 3333381020 and f["DescSituacion"] == "ENTREGADO"
    assert f["NumeroLiquidacion"] == 2 and f["id"] == -37
    assert f["CONCATENADO"] == "7870740-15" and f["FECHA DE CIERRE"] == datetime(2026, 9, 14)
    assert f["MOROCIDAD"] == "Mora 3" and f["CAMPAÑA DE TRABAJO"] == 19 and f["REFERENCIA DE PAGO"] == "7870740-3"
    assert f["Direccion"] == "TULE No 79 COLONIA ALAMEDA CP 45679" and f["ColumnaNueva"] == "x"


def test_revision_marca_solo_celdas_calculadas_vacias():
    wb = _generar()
    ws = wb["BASE"]
    enc = [c.value.strip() for c in ws[1]]
    celda = {e: ws.cell(3, i + 1) for i, e in enumerate(enc)}  # fila 3: ZONA 999, CP sin catálogo
    assert celda["REGION"].value is None and celda["REGION"].fill.fgColor.rgb == "FFFFFF00"
    assert celda["Municipio / Poblacion"].fill.fgColor.rgb == "FFFFFF00"
    assert celda["Nombre"].fill.fgColor.rgb in (None, "00000000")  # dato de la cartera: sin marca
    assert celda["CONCATENADO"].fill.fgColor.rgb in (None, "00000000")  # calculado y lleno: sin marca
    assert wb["Revision"].max_row == 2
