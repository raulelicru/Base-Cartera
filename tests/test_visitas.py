import io
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook

import procesamiento as proc
from tests.test_procesamiento import _estructura_xlsx

CATALOGO = pd.DataFrame(
    {"Cp": ["07140"], "Municipio": ["Gustavo A. Madero"], "Estado": ["Ciudad de México"], "Zona": ["Urbano"]}
)


def _cartera_xlsx(con_extras=True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BASE"
    enc = ["ZONA", "NoDama", "Direccion", "Referencia", "AnioSaldo", "CampaniaSaldo", "DigitoVerificador"]
    if con_extras:
        # Nombres como vienen en la cartera real (TelefonoCelular está en la columna W)
        enc += ["NOMBRE", "SaldoDama", "DescSituacionCie", "ImporteOtro"] + [f"X{i}" for i in range(11)] + ["TelefonoCelular"]
    ws.append(enc)
    filas = [
        [101, 7241893, "JOSE ALFREDO JIMENEZ No 11 Mza 37 Lt 11 COLONIA FORESTAL CP 7140",
         "JUVENTINO ROSAS Y JOS ALFREDO JIMENEZ", 2025, 15, 2],
        [None] * 7,
        [102, 2312748, "JOSE MIJICA Mza 51 Lt 473 LA FORESTAL CP 7140", "MIMI DERBA Y JORGE NEGRETE", 2025, 18, 3],
    ]
    extras = [["CARMEN DELIA HERNANDEZ SANTOS", 459, "ACTIVA", 1] + [None] * 11 + [5512345678], [None] * 16,
              ["KARINA TAPIA RAMIREZ", 738, "SUSPENDIDA", 2] + [None] * 11 + [None]]
    for f, e in zip(filas, extras):
        ws.append(f + (e if con_extras else []))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _resultado(contenido):
    cartera = proc.leer_cartera(contenido, "Cartera_Campaña_19.xlsx")
    res = proc.generar_base(cartera, proc.leer_estructura(_estructura_xlsx()), CATALOGO, 19)
    return cartera, res


def test_tabla_visitas_completa():
    cartera, res = _resultado(_cartera_xlsx())
    assert proc.columnas_visitas_faltantes(cartera) == []
    t = proc.tabla_visitas(res, cartera, date(2026, 9, 24))
    assert list(t.columns) == proc.COLUMNAS_VISITAS and len(t) == 2
    f = t.iloc[0].to_dict()
    assert f == {
        "FECHA DE ASIGNACION": datetime(2026, 9, 24),
        "ZONA": 101,
        "ASIGNACION": None,
        "NoDama": 7241893,
        "DIGITO VERIFICADOR": 2,
        "NOMBRE": "CARMEN DELIA HERNANDEZ SANTOS",
        "DIRECCION": "Jose Alfredo Jimenez No. 11 Mz 37 L- 11",
        "COLONIA": "FORESTAL",
        "CP Extraido": 7140,
        "LOCALIDAD": "GUSTAVO A MADERO",
        "REFERENCIA": "JUVENTINO ROSAS Y JOS ALFREDO JIMENEZ",
        "TEMPORALIDAD": "Mora 3",
        "CAMPANA": 202515,
        "IMPORTE NETO FACTURA": 459,
        "TELEFONO CELULAR": 5512345678,
        "DescSituacionCie": "ACTIVA",
    }
    g = t.iloc[1]
    assert g["COLONIA"] == "LA FORESTAL" and g["TEMPORALIDAD"] == "Mora 1" and g["NOMBRE"] == "KARINA TAPIA RAMIREZ"
    # Todas las columnas llenas excepto ASIGNACION (y el teléfono que no venía en la cartera)
    vacias = {c for c in proc.COLUMNAS_VISITAS if t[c].isna().all()}
    assert vacias == {"ASIGNACION"}


def test_telefono_en_columna_w():
    cartera, _ = _resultado(_cartera_xlsx())
    assert list(cartera.columns).index("TelefonoCelular") == 22  # columna W


def test_visitas_sin_columnas_extra_avisa():
    cartera, res = _resultado(_cartera_xlsx(con_extras=False))
    assert proc.columnas_visitas_faltantes(cartera) == list(proc.CAMPOS_CARTERA_VISITAS)
    t = proc.tabla_visitas(res, cartera, date(2026, 9, 24))
    assert t["NOMBRE"].isna().all() and t["NoDama"].notna().all()


def test_exportar_visitas_usa_plantilla():
    cartera, res = _resultado(_cartera_xlsx())
    t = proc.tabla_visitas(res, cartera, date(2026, 9, 24))
    wb = load_workbook(io.BytesIO(proc.exportar_visitas(t, date(2026, 9, 24))))
    ws = wb.active
    assert ws.title == "VISITAS 24 SEP"
    assert [c.value for c in ws[1]] == proc.COLUMNAS_VISITAS
    assert ws.max_row == 3
    assert ws["A2"].value == datetime(2026, 9, 24) and ws["A2"].number_format == "d-mmm"
    assert ws["C2"].value is None and ws["F3"].value == "KARINA TAPIA RAMIREZ"
    # Mismos colores que la base de visitas de ejemplo: COLONIA y LOCALIDAD en amarillo
    assert ws["H2"].fill.fgColor.rgb == "FFFFFF00" and ws["J3"].fill.fgColor.rgb == "FFFFFF00"
    assert ws["A1"].fill.fgColor.theme == 9
    assert ws.column_dimensions["G"].width > 80


def test_visitas_rapido_con_muchas_filas():
    import time

    cartera, res = _resultado(_cartera_xlsx())
    grande = proc.Resultado(
        base=pd.concat([res.base] * 15000, ignore_index=True).assign(_fila_excel=None),
        resumen=res.resumen, advertencias=[],
    )
    inicio = time.perf_counter()
    t = proc.tabla_visitas(grande, None, date(2026, 9, 24))
    proc.exportar_visitas(t, date(2026, 9, 24))
    assert len(t) == 30000
    assert time.perf_counter() - inicio < 30
