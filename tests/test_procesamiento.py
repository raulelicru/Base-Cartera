import io
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import procesamiento as proc


def _estructura_xlsx() -> bytes:
    wb = Workbook()
    z = wb.active
    z.title = "Base de Zonas"
    z.append(["ZONA", "REGION", "DIVISION", "RUTA", "NO. COBRADOR"])
    z.append([101, "CENTRO", "DIV 1", 5, 9001])
    z.append([102, "NORTE", "DIV 2", 7, 9002])

    cal = wb.create_sheet("Calendario de Cierre")
    cal.cell(1, 1, "Campaña de Trabajo 2")
    cal.cell(2, 1, "Ruta"); cal.cell(2, 2, "Fecha")
    cal.cell(1, 3, "Campaña de Trabajo 19")
    cal.cell(2, 3, "Ruta"); cal.cell(2, 4, "Fecha Inicio"); cal.cell(2, 5, "Fecha Cierre")
    cal.cell(1, 6, "Campaña de Trabajo 20")
    cal.cell(2, 6, "Ruta"); cal.cell(2, 7, "Fecha Inicio"); cal.cell(2, 8, "Fecha Cierre")
    cal.cell(3, 3, 5); cal.cell(3, 4, datetime(2026, 9, 1)); cal.cell(3, 5, datetime(2026, 9, 14))
    cal.cell(3, 6, 5); cal.cell(3, 7, datetime(2026, 9, 15)); cal.cell(3, 8, datetime(2026, 9, 28))

    cam = wb.create_sheet("Campaña de Trabajo")
    cam.cell(1, 1, "Campaña de Trabajo 19")
    cam.cell(2, 1, "CAMPAÑA"); cam.cell(2, 2, "Mora")
    for i, (c, m) in enumerate([(18, 1), (17, 1), (16, 2), (15, 3), (26, 3)], start=3):
        cam.cell(i, 1, c); cam.cell(i, 2, m)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _cartera():
    return pd.DataFrame({
        "ZONA": [101, 102, 999],
        "NoDama": [7601974, 7601975, 7601976],
        "Direccion": [
            "PV TECOLOTE No 37 Int a COL CENTRO CP 5050",
            "HECTOR ALVARADO Mza I Lt 12 BARRIO SAN JUAN CP 50000",
            "TENANGO DEL VALLE No 0 Mza 167 Lt 18 LA LOMA CP 00000",
        ],
        "Referencia": ["ENTRE A Y B", "", None],
        "AnioSaldo": [2026, 2026, 2026],
        "CampaniaSaldo": [15, 18, 3],
        "DigitoVerificador": [4, 5, 6],
        "REGION": [None] * 3,
    })


def test_extraer_cp():
    assert proc.extraer_cp("CALLE X CP 5050") == "05050"
    assert proc.extraer_cp("CALLE X CP00000") is None
    assert proc.extraer_cp("CALLE X") is None


@pytest.mark.parametrize("entrada,calle", [
    ("PV TECOLOTE No 37 Int a COL CENTRO CP 5050", "PV TECOLOTE NO. 37 INTERIOR A"),
    ("HECTOR ALVARADO Mza I Lt 12 LOS PINOS CP 50000", "HECTOR ALVARADO MZ I L- 12"),
    ("TENANGO DEL VALLE No 0 Mza 167 Lt 18 LA LOMA CP 50000", "TENANGO DEL VALLE NO. 0 MZ 167 L- 18"),
])
def test_direccion_calle(entrada, calle):
    assert proc.segmentar_direccion(entrada).direccion_calle() == calle


def test_colonia():
    assert proc.segmentar_direccion("PV TECOLOTE No 37 Int a COL CENTRO CP 5050").colonia == "COL CENTRO"


def test_inferir_campania():
    assert proc.inferir_campania("Cartera_Campaña_19.xlsx") == 19
    assert proc.inferir_campania("Cartera_Campana_7.xlsx") == 7


def test_catalogo_txt_sepomex():
    txt = (
        "El Catálogo Nacional de Códigos Postales...\n"
        "d_codigo|d_asenta|d_tipo_asenta|D_mnpio|d_estado|d_ciudad|d_zona\n"
        "05050|Centro|Colonia|Cuajimalpa|Ciudad de México|CDMX|Urbano\n"
        "50000|San Juan|Barrio|Toluca|México|Toluca|\n"
    ).encode("latin-1")
    cat = proc.leer_catalogo_cp(txt, "CPdescarga.txt")
    d = cat.set_index("Cp").to_dict("index")
    assert d["05050"]["Zona"] == "Urbano"
    assert d["50000"]["Zona"] == "Semiurbano"  # inferido de 'Barrio'
    assert d["50000"]["Municipio"] == "Toluca"


def test_generar_base_completa():
    estructura = proc.leer_estructura(_estructura_xlsx())
    assert estructura.campanias_disponibles() == [19]
    catalogo = pd.DataFrame({"Cp": ["05050"], "Municipio": ["Cuajimalpa"], "Estado": ["CDMX"], "Zona": ["Urbano"]})
    cartera = proc._normalizar_columnas_cartera(_cartera())
    res = proc.generar_base(cartera, estructura, catalogo, 19)
    b = res.base
    assert list(b.columns[:22]) == proc.COLUMNAS_SALIDA

    f0 = b.iloc[0]
    assert f0["REGION"] == "CENTRO" and f0["RUTA"] == 5 and f0["ID COBRADOR"] == 9001
    assert f0["Concatenado"] == "7601974-15"
    assert f0["Referencia de Pago"] == "7601974-4"
    assert f0["Fecha de cierre"] == datetime(2026, 9, 14)
    assert f0["Morosidad"] == 3
    assert f0["Campaña de trabajo"] == 19
    assert f0["Municipio / Poblacion"] == "Cuajimalpa"
    assert not f0["requiere_revision"]

    f1 = b.iloc[1]  # ruta 7 sin fecha, CP no en catálogo
    assert f1["Morosidad"] == 1 and f1["Fecha de cierre"] is None and f1["requiere_revision"]

    f2 = b.iloc[2]  # zona inexistente, CP 00000, campaña 3 sin mora
    assert f2["Cp"] is None and f2["REGION"] is None and f2["Morosidad"] is None
    assert "ZONA no encontrada" in f2["motivo_revision"]

    r = res.resumen
    assert r["Total de filas procesadas"] == 3
    assert r["Zona encontrada"] == 2 and r["CP no identificado"] == 1
    assert r["Filas que requieren revisión"] == 2

    xlsx = load_workbook(io.BytesIO(proc.exportar_excel(res, 19)))
    ws = xlsx["Base de Gestion"]
    assert ws.cell(2, 1).fill.fgColor.rgb in (None, "00000000")
    assert ws.cell(3, 1).fill.fgColor.rgb.endswith("FFFF00")
    assert set(xlsx.sheetnames) == {"Base de Gestion", "Revision", "Resumen"}


def test_campania_inexistente():
    estructura = proc.leer_estructura(_estructura_xlsx())
    with pytest.raises(ValueError, match="Campaña de Trabajo 25"):
        proc.generar_base(proc._normalizar_columnas_cartera(_cartera()), estructura, None, 25)


def test_columna_zona_duplicada():
    df = _cartera().assign(Zona=[1, 2, 3])
    out = proc._normalizar_columnas_cartera(df)
    assert out["ZONA"].tolist() == [101, 102, 999]


def _estructura_apilada() -> bytes:
    """Bloques 19 arriba y 21 debajo (otro renglón de bloques), con encabezados en mayúsculas."""
    wb = Workbook()
    z = wb.active
    z.title = "Base de Zonas"
    z.append(["ZONA", "REGION", "DIVISION", "RUTA", "NO. COBRADOR"])
    z.append([101, "CENTRO", "DIV 1", 5, 9001])
    cal = wb.create_sheet("Calendario de Cierre")
    cal.cell(1, 1, "Campaña de Trabajo 19")
    cal.cell(2, 1, "Ruta"); cal.cell(2, 2, "Fecha Inicio"); cal.cell(2, 3, "Fecha Cierre")
    cal.cell(3, 1, 5); cal.cell(3, 3, datetime(2026, 9, 14))
    cal.cell(30, 1, "CAMPAÑA DE TRABAJO 21")
    cal.cell(31, 1, "Ruta"); cal.cell(31, 2, "Fecha Inicio"); cal.cell(31, 3, "Fecha Cierre")
    cal.cell(32, 1, 5); cal.cell(32, 3, datetime(2026, 10, 12))
    cam = wb.create_sheet("Campaña de Trabajo")
    cam.cell(1, 1, "Campaña de Trabajo 19"); cam.cell(2, 1, "CAMPAÑA"); cam.cell(2, 2, "Mora")
    cam.cell(3, 1, 15); cam.cell(3, 2, 3)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def test_campania_en_bloque_inferior_y_diagnostico():
    e = proc.leer_estructura(_estructura_apilada())
    assert proc.tabla_calendario(e, 21) == {"5": datetime(2026, 10, 12)}
    assert proc.tabla_calendario(e, 19) == {"5": datetime(2026, 9, 14)}  # no se mezcla con el bloque 21
    assert e.campanias_disponibles() == [19]
    msg = e.diagnostico_campania(21)
    assert "Campaña de Trabajo" in msg and "Calendario de Cierre" not in msg.split("(")[0]
    assert e.diagnostico_campania(19) is None
