import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import procesamiento as proc
from tests.test_procesamiento import _estructura_xlsx

VERDE = PatternFill(start_color="FF00B050", end_color="FF00B050", fill_type="solid")
NARANJA = PatternFill(start_color="FFFFC000", end_color="FFFFC000", fill_type="solid")


def _cartera_con_formato() -> bytes:
    wb = Workbook()
    wb.active.title = "Notas"
    wb["Notas"]["A1"] = "hoja que no se toca"
    ws = wb.create_sheet("BASE")
    enc = ["ZONA", "NoDama", "Direccion", "Referencia", "AnioSaldo", "CampaniaSaldo", "DigitoVerificador",
           "REGION", "RUTA", "DIVISION", "ID COBRADOR", "Concatenado", "Fecha de cierre", "Morosidad",
           "Campaña de trabajo", "Referencia de Pago", "Extra usuario"]
    ws.append(enc)
    for c in ws[1]:
        c.fill = VERDE
        c.font = Font(bold=True, color="FFFFFFFF")
    ws.append([101, 7601974, "PV TECOLOTE No 37 Int a COL CENTRO CP 5050", "ENTRE A Y B", 2026, 15, 4] + [None] * 9 + ["x1"])
    ws.append([None] * 17)  # fila en blanco en medio
    ws.append([999, 7601976, "SIN NUMERO CP 00000", None, 2026, 3, 6] + [None] * 9 + ["x3"])
    ws["C2"].fill = NARANJA  # color puesto por el usuario en una celda de datos
    ws.column_dimensions["C"].width = 55
    ws.auto_filter.ref = "A1:Q4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generar(contenido, **kw):
    estructura = proc.leer_estructura(_estructura_xlsx())
    catalogo = pd.DataFrame({"Cp": ["05050"], "Municipio": ["Cuajimalpa"], "Estado": ["CDMX"], "Zona": ["Urbano"]})
    cartera = proc.leer_cartera(contenido, "Cartera_Campaña_19.xlsx")
    res = proc.generar_base(cartera, estructura, catalogo, 19)
    salida = proc.exportar_excel_original(contenido, "Cartera_Campaña_19.xlsx", res, 19, **kw)
    return cartera, res, load_workbook(io.BytesIO(salida))


def test_leer_cartera_guarda_fila_original():
    cartera = proc.leer_cartera(_cartera_con_formato(), "Cartera_Campaña_19.xlsx")
    assert cartera["_fila_excel"].tolist() == [2, 4]  # la fila vacía (3) se omite
    assert cartera["NoDama"].tolist() == [7601974, 7601976]


def test_salida_conserva_cartera_y_anexa_columnas():
    _, _, wb = _generar(_cartera_con_formato(), color_sistema="#FFC7CE")
    assert wb.sheetnames[:2] == ["Notas", "BASE"] and "Revision" in wb.sheetnames and "Resumen" in wb.sheetnames
    assert wb["Notas"]["A1"].value == "hoja que no se toca"
    ws = wb["BASE"]
    enc = [c.value for c in ws[1]]

    # Columnas originales: mismo orden, mismos colores, mismos datos
    assert enc[:17][0] == "ZONA" and enc[16] == "Extra usuario"
    assert all(ws.cell(1, c).fill.fgColor.rgb == "FF00B050" for c in range(1, 18))
    assert ws["C2"].fill.fgColor.rgb == "FFFFC000"
    assert ws.column_dimensions["C"].width == 55
    assert ws["Q2"].value == "x1" and ws["Q4"].value == "x3"
    assert all(ws.cell(3, c).value is None for c in range(1, 18))

    # Columnas vacías que ya existían: se llenan en su lugar sin cambiar su color
    assert ws["H2"].value == "CENTRO" and ws["I2"].value == 5 and ws["K2"].value == 9001
    assert ws["L2"].value == "7601974-15" and ws["M2"].value == datetime(2026, 9, 14)
    assert ws["N2"].value == 3 and ws["O2"].value == 19 and ws["P2"].value == "7601974-4"
    assert ws["H2"].fill.fgColor.rgb in (None, "00000000")

    # Columnas anexadas al final, con el color elegido
    anexadas = enc[17:]
    assert anexadas == ["Direccion Calle", "Colonia", "Municipio / Poblacion", "Cp", "Estado",
                        "Zona (Urbano/Rural)", proc.COLUMNA_MOTIVO]
    assert all(ws.cell(1, c).fill.fgColor.rgb == "FFFFC7CE" for c in range(18, 25))
    assert ws.cell(2, 18).value == "Pv Tecolote No. 37 Interior a" and ws.cell(2, 21).value == "05050"
    assert ws.cell(2, 18).fill.fgColor.rgb == "FFFFC7CE"
    # Fila que requiere revisión: celdas anexadas en amarillo y con motivo
    assert ws.cell(4, 18).fill.fgColor.rgb == "FFFFFF00"
    assert "ZONA no encontrada" in ws.cell(4, 24).value
    assert ws.cell(2, 24).value is None
    assert ws.auto_filter.ref == "A1:X4"


def test_colorear_columnas_llenadas():
    _, _, wb = _generar(_cartera_con_formato(), color_sistema="#BDD7EE", colorear_llenadas=True)
    ws = wb["BASE"]
    assert ws["H1"].fill.fgColor.rgb == "FFBDD7EE" and ws["H2"].fill.fgColor.rgb == "FFBDD7EE"
    assert ws["A1"].fill.fgColor.rgb == "FF00B050"  # columnas de datos originales intactas
